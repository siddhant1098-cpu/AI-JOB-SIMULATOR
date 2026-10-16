import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image, ImageDraw, ImageFont

# ── Color Palette (Minimal Pastel & Dark Slate) ──────────────────────────────
BG_COLOR = RGBColor(250, 248, 245)       # Alabaster Cream (#FAF8F5)
CARD_BG = RGBColor(255, 255, 255)        # Pure White
CARD_BORDER = RGBColor(226, 222, 213)    # Soft Stone Border (#E2DED5)
TEXT_PRIMARY = RGBColor(30, 41, 59)      # Slate 800 (#1E293B)
TEXT_MUTED = RGBColor(100, 116, 139)     # Slate 500 (#64748B)

# Pastel Accents
ACCENT_LILAC = RGBColor(124, 107, 181)   # Wisteria / Lilac (#7C6BB5)
LILAC_BG = RGBColor(242, 237, 253)       # Soft Lavender (#F2EDFD)
ACCENT_MINT = RGBColor(45, 106, 79)      # Soft Mint (#2D6A4F)
MINT_BG = RGBColor(238, 247, 242)        # Soft Sage (#EEF7F2)
ACCENT_PEACH = RGBColor(194, 91, 50)     # Soft Peach / Amber (#C25B32)
PEACH_BG = RGBColor(254, 247, 236)       # Soft Apricot (#FEF7EC)
ACCENT_ROSE = RGBColor(190, 40, 60)      # Rose (#BE283C)
ROSE_BG = RGBColor(253, 240, 237)        # Soft Rose (#FDF0ED)

os.makedirs("presentation_assets", exist_ok=True)

# ── Generate Clean UI Mockup Graphics ─────────────────────────────────────────
def create_mockup(filename, title, subtitle, badge_text, items, alert=None):
    img = Image.new("RGBA", (900, 520), (250, 248, 245, 255))
    draw = ImageDraw.Draw(img)

    # Card background with shadow simulation
    draw.rounded_rectangle([22, 22, 882, 502], radius=24, fill=(235, 230, 222, 100))
    draw.rounded_rectangle([20, 20, 880, 500], radius=24, fill=(255, 255, 255, 255), outline=(226, 222, 213, 255), width=2)

    # Header bar
    draw.rectangle([20, 20, 880, 85], fill=(248, 246, 242, 255))
    draw.line([20, 85, 880, 85], fill=(226, 222, 213, 255), width=1)

    # Header details
    draw.text((45, 42), title, fill=(30, 41, 59))
    draw.text((45, 62), subtitle, fill=(100, 116, 139))
    
    # Badge
    draw.rounded_rectangle([720, 38, 850, 68], radius=15, fill=(242, 237, 253, 255), outline=(221, 213, 247, 255), width=1)
    draw.text((735, 47), badge_text, fill=(124, 107, 181))

    # Content cards
    y = 105
    for item_title, item_desc, icon_type in items:
        bg_col = (246, 243, 254, 255) if icon_type == "purple" else (238, 247, 242, 255) if icon_type == "green" else (254, 247, 236, 255)
        border_col = (226, 219, 250, 255) if icon_type == "purple" else (205, 229, 215, 255) if icon_type == "green" else (248, 225, 194, 255)
        
        draw.rounded_rectangle([45, y, 855, y + 80], radius=16, fill=bg_col, outline=border_col, width=1)
        draw.text((65, y + 18), item_title, fill=(30, 41, 59))
        draw.text((65, y + 44), item_desc, fill=(71, 85, 105))
        y += 95

    # Alert banner if present
    if alert:
        draw.rounded_rectangle([45, 415, 855, 475], radius=16, fill=(253, 240, 237, 255), outline=(248, 210, 202, 255), width=1)
        draw.text((65, 435), alert, fill=(190, 40, 60))

    img.save(filename)

# Mockup 1: Simulation Screen
create_mockup(
    "presentation_assets/screen_simulation.png",
    "AI Job Simulator  •  Operations Manager",
    "Domain: Event Management  •  Level: Balanced  •  Time: 14:28",
    "🛡️ Strict Mode",
    [
        ("Scenario Situation 1", "Your lead AV equipment supplier calls 48 hours prior to an international summit to cancel. How do you respond?", "green"),
        ("Live Spoken Response (Speech-To-Text Active)", "\"I would immediately contact pre-approved emergency backup vendors, inform core organizers, and adjust production schedules...\"", "purple"),
        ("Empathetic Mentor Feedback", "Thoughtful response! I loved how you stayed composed and prioritized the immediate attendee experience. Rated 8.8/10.", "green")
    ]
)

# Mockup 2: Anti-Cheat Camera Alert
create_mockup(
    "presentation_assets/screen_proctor.png",
    "AI Job Simulator  •  On-Device Vision Proctor",
    "Real-Time Surveillance  •  MediaPipe Tasks Vision WASM",
    "📱 Phone Check",
    [
        ("MediaPipe BlazeFace Presence Sensor", "Continuous 500ms facial verification running 100% on-device. Zero video frames uploaded.", "purple"),
        ("MediaPipe EfficientDet-Lite0 Object Recognition", "Trained on COCO dataset. Automatically flags smartphones, tablets, or unauthorized devices in frame.", "purple")
    ],
    alert="⚠️ Strict Mode Alert: Mobile device detected in webcam frame! Assessment paused until phone is removed."
)

# ── Build PowerPoint Presentation ────────────────────────────────────────────
prs = Presentation()
prs.slide_width = Inches(13.333)   # 16:9 Widescreen
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

def set_slide_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR

def add_header(slide, title_text, category_text="AI JOB SIMULATOR"):
    # Category pill
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8), Inches(0.4))
    tf = cat_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = category_text.upper()
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_LILAC

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.5), Inches(0.8))
    tf2 = title_box.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = title_text
    p2.font.size = Pt(26)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_PRIMARY

# ── Slide 1: Title Slide ──────────────────────────────────────────────────────
s1 = prs.slides.add_slide(blank_layout)
set_slide_background(s1)

# Main Title Card
main_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.2), Inches(10.9), Inches(5.1))
main_card.fill.solid()
main_card.fill.fore_color.rgb = CARD_BG
main_card.line.color.rgb = CARD_BORDER
main_card.line.width = Pt(1.5)

tf = main_card.text_frame
tf.word_wrap = True
tf.margin_left = Inches(0.8)
tf.margin_top = Inches(0.8)

p = tf.paragraphs[0]
p.text = "🌿  AI JOB SIMULATOR"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_LILAC

p2 = tf.add_paragraph()
p2.text = "Human-Centered Workplace Simulation &"
p2.font.size = Pt(38)
p2.font.bold = True
p2.font.color.rgb = TEXT_PRIMARY
p2.space_before = Pt(10)

p3 = tf.add_paragraph()
p3.text = "On-Device Anti-Cheat Proctoring"
p3.font.size = Pt(38)
p3.font.bold = True
p3.font.color.rgb = ACCENT_LILAC

p4 = tf.add_paragraph()
p4.text = "Real-time conversational interviews, empathetic AI mentoring, and privacy-preserving computer vision."
p4.font.size = Pt(15)
p4.font.color.rgb = TEXT_MUTED
p4.space_before = Pt(18)

p5 = tf.add_paragraph()
p5.text = "Speakers: 4-Member Team   •   Time Limit: 8 Minutes   •   Stack: React 19, FastAPI, Google Gemini, MediaPipe"
p5.font.size = Pt(12)
p5.font.bold = True
p5.font.color.rgb = ACCENT_MINT
p5.space_before = Pt(30)

# ── Slide 2: Problem & Solution ───────────────────────────────────────────────
s2 = prs.slides.add_slide(blank_layout)
set_slide_background(s2)
add_header(s2, "The Assessment Dilemma: Robotic Tests vs. Real-World Readiness")

# Problem Card
c1 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9))
c1.fill.solid()
c1.fill.fore_color.rgb = ROSE_BG
c1.line.color.rgb = RGBColor(248, 210, 202)
c1.line.width = Pt(1.5)
tf = c1.text_frame
tf.margin_left = Inches(0.5)
tf.margin_top = Inches(0.5)
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "❌  THE OLD WAY"
p.font.bold = True
p.font.size = Pt(14)
p.font.color.rgb = ACCENT_ROSE

bullet_pts = [
    ("Cold, Robotic Interfaces: ", "Static multiple-choice questionnaires that do not test how people speak or collaborate."),
    ("Harsh & Rigid AI Grading: ", "Pedantic keyword matching penalizes everyday human explanations that possess solid common sense."),
    ("Remote Integrity Vulnerability: ", "Exams are vulnerable to mobile phones, unauthorized second persons, and copy-pasting.")
]
for title, desc in bullet_pts:
    p = tf.add_paragraph()
    p.text = "• " + title + desc
    p.font.size = Pt(13)
    p.font.color.rgb = TEXT_PRIMARY
    p.space_before = Pt(16)

# Solution Card
c2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.9))
c2.fill.solid()
c2.fill.fore_color.rgb = MINT_BG
c2.line.color.rgb = RGBColor(205, 229, 215)
c2.line.width = Pt(1.5)
tf = c2.text_frame
tf.margin_left = Inches(0.5)
tf.margin_top = Inches(0.5)
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "✓  OUR SOLUTION: AI JOB SIMULATOR"
p.font.bold = True
p.font.size = Pt(14)
p.font.color.rgb = ACCENT_MINT

sol_pts = [
    ("Spoken Human Dialogue: ", "Candidates listen to a natural human voice and speak their answers naturally aloud."),
    ("Empathetic Practical Evaluation: ", "Grades practical problem solving, empathy, and composure with generous, mentor-like coaching."),
    ("100% On-Device Anti-Cheat: ", "Real-time MediaPipe Vision detects phones, absence, and blocks clipboard shortcuts with zero cloud video upload.")
]
for title, desc in sol_pts:
    p = tf.add_paragraph()
    p.text = "• " + title + desc
    p.font.size = Pt(13)
    p.font.color.rgb = TEXT_PRIMARY
    p.space_before = Pt(16)

# ── Slide 3: Architecture & Tech Stack ────────────────────────────────────────
s3 = prs.slides.add_slide(blank_layout)
set_slide_background(s3)
add_header(s3, "Full-Stack System Architecture & Technology Foundation")

stack_boxes = [
    ("Frontend Layer", "React 19 + Vite + Tailwind v4", "Minimal pastel UI, client-side session routing, Recharts visual analytics, and reactive question runner.", LILAC_BG, ACCENT_LILAC),
    ("AI & Reasoning Engine", "Google Gemini 3.5 Flash", "Sub-second workplace scenario generation, multi-criteria answer evaluation, and json-repair fault tolerance.", MINT_BG, ACCENT_MINT),
    ("Vision Proctoring", "MediaPipe Tasks Vision (WASM)", "Client-side BlazeFace presence check and EfficientDet-Lite0 mobile phone detection running 100% in browser.", PEACH_BG, ACCENT_PEACH),
    ("Backend API Layer", "FastAPI (Python) + Uvicorn", "High-throughput asynchronous REST API, Pydantic v2 schemas, and zero-database stateless architecture.", LILAC_BG, ACCENT_LILAC)
]

for idx, (title, tech, desc, bg, acc) in enumerate(stack_boxes):
    col = idx % 2
    row = idx // 2
    x = Inches(0.8 + col * 6.0)
    y = Inches(1.8 + row * 2.5)

    box = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.6), Inches(2.2))
    box.fill.solid()
    box.fill.fore_color.rgb = bg
    box.line.color.rgb = acc
    box.line.width = Pt(1)

    tf = box.text_frame
    tf.margin_left = Inches(0.4)
    tf.margin_top = Inches(0.3)
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = title.upper()
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = acc

    p2 = tf.add_paragraph()
    p2.text = tech
    p2.font.size = Pt(16)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_PRIMARY
    p2.space_before = Pt(4)

    p3 = tf.add_paragraph()
    p3.text = desc
    p3.font.size = Pt(11)
    p3.font.color.rgb = TEXT_MUTED
    p3.space_before = Pt(6)

# ── Slide 4: Conversational AI & Empathetic Evaluation ────────────────────────
s4 = prs.slides.add_slide(blank_layout)
set_slide_background(s4)
add_header(s4, "Conversational Speech & Fair, Empathetic Human Evaluation")

# Left Column: Features
c_left = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9))
c_left.fill.solid()
c_left.fill.fore_color.rgb = CARD_BG
c_left.line.color.rgb = CARD_BORDER
c_left.line.width = Pt(1.5)
tf = c_left.text_frame
tf.margin_left = Inches(0.4)
tf.margin_top = Inches(0.4)
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "NATURAL SPEECH & FAIR MARKING"
p.font.bold = True
p.font.size = Pt(13)
p.font.color.rgb = ACCENT_LILAC

points = [
    ("Warm Spoken Tone (0.95x): ", "Auto-selects top natural human voices in Chrome/Edge. Questions flow conversationally like a senior colleague over coffee."),
    ("Respects Human Speech: ", "Spoken answers are transcribed live. The evaluation ignores filler words and values common sense, empathy, and practical initiative."),
    ("Generous Scoring Bands: ", "Thoughtful practical answers score 7.5 to 9.5 / 10 instead of harsh penalties for missing textbook terminology."),
    ("Supportive Mentoring: ", "Every evaluation delivers genuine praise, an ideal framing from a senior peer, and gentle growth steps.")
]
for title, desc in points:
    p = tf.add_paragraph()
    p.text = "• " + title + desc
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_PRIMARY
    p.space_before = Pt(14)

# Right Column: Screenshot Mockup
s4.shapes.add_picture("presentation_assets/screen_simulation.png", Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.9))

# ── Slide 5: Strict Anti-Cheat & On-Device Computer Vision ────────────────────
s5 = prs.slides.add_slide(blank_layout)
set_slide_background(s5)
add_header(s5, "Strict Anti-Cheat Proctoring: On-Device Vision & Clipboard Lock")

# Left Column: Screenshot Mockup
s5.shapes.add_picture("presentation_assets/screen_proctor.png", Inches(0.8), Inches(1.8), Inches(5.7), Inches(4.9))

# Right Column: Tech Details
c_right = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.9))
c_right.fill.solid()
c_right.fill.fore_color.rgb = CARD_BG
c_right.line.color.rgb = CARD_BORDER
c_right.line.width = Pt(1.5)
tf = c_right.text_frame
tf.margin_left = Inches(0.4)
tf.margin_top = Inches(0.4)
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "MULTI-LAYER PROCTORING DEFENSE"
p.font.bold = True
p.font.size = Pt(13)
p.font.color.rgb = ACCENT_PEACH

sec_points = [
    ("📱 AI Mobile Phone Detection: ", "MediaPipe EfficientDet-Lite0 scans every video frame for smartphones. Detecting a phone immediately blurs the screen and pauses the exam."),
    ("👤 Presence & Multi-Person Lock: ", "MediaPipe BlazeFace verifies single-candidate presence every 500ms. Leaving the frame or a second face triggers an instant blur."),
    ("🚫 Clipboard & Text Lock: ", "Text selection is disabled. Copy, cut, and paste events are intercepted and blocked with real-time warning toasts."),
    ("📑 Tab-Switch & Inactivity Tracking: ", "Monitors document visibility. Navigating away to another browser tab immediately blurs the test and logs a violation.")
]
for title, desc in sec_points:
    p = tf.add_paragraph()
    p.text = "• " + title + desc
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_PRIMARY
    p.space_before = Pt(14)

# ── Slide 6: Performance Report & Proctor Audit ───────────────────────────────
s6 = prs.slides.add_slide(blank_layout)
set_slide_background(s6)
add_header(s6, "Comprehensive Debrief: Visual Analytics & Proctor Audit")

report_boxes = [
    ("Balanced Evaluation Rubric", "Exact 40 / 30 / 30 Math", "Evaluates Practical Solution (40%), Composure & Problem Solving (30%), and Communication (30%) computed with exact code arithmetic.", MINT_BG, ACCENT_MINT),
    ("Interactive Visual Analytics", "Recharts Competency Radar", "Radar profile and performance bar charts provide candidate and recruiter with an instant graphical overview of core strengths.", LILAC_BG, ACCENT_LILAC),
    ("Empowering Career Mentoring", "Actionable Next Steps", "Identifies 3 concrete strengths and 2 gentle growth areas, plus 3 tailored next steps to prepare candidates for real interviews.", PEACH_BG, ACCENT_PEACH),
    ("Strict Proctor Integrity Audit", "Full Transparent Audit Trail", "Reports exact counts of Mobile Phone sightings, Clipboard paste attempts, Tab switches, and Face absences with verified clean badges.", ROSE_BG, ACCENT_ROSE)
]

for idx, (title, sub, desc, bg, acc) in enumerate(report_boxes):
    col = idx % 2
    row = idx // 2
    x = Inches(0.8 + col * 6.0)
    y = Inches(1.8 + row * 2.5)

    box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.6), Inches(2.2))
    box.fill.solid()
    box.fill.fore_color.rgb = bg
    box.line.color.rgb = acc
    box.line.width = Pt(1)

    tf = box.text_frame
    tf.margin_left = Inches(0.4)
    tf.margin_top = Inches(0.3)
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = TEXT_PRIMARY

    p2 = tf.add_paragraph()
    p2.text = sub
    p2.font.size = Pt(11)
    p2.font.bold = True
    p2.font.color.rgb = acc
    p2.space_before = Pt(2)

    p3 = tf.add_paragraph()
    p3.text = desc
    p3.font.size = Pt(11)
    p3.font.color.rgb = TEXT_MUTED
    p3.space_before = Pt(6)

# ── Slide 7: Zero-Database Architecture & 1-Click Deployment ──────────────────
s7 = prs.slides.add_slide(blank_layout)
set_slide_background(s7)
add_header(s7, "Engineering Excellence: Zero-Database Resilience & 1-Click Launch")

eng_boxes = [
    ("Zero-Database Stateless Architecture", "Client-Side LocalStorage Engine", "No user database or login barriers. Candidate sessions, question states, audio transcripts, and integrity logs persist locally in browser localStorage.", LILAC_BG, ACCENT_LILAC),
    ("Mid-Test Refresh Resume", "Timestamp-Based Recovery", "Accidental browser refreshes or disconnects show a 1-click 'Resume Session' screen, re-requesting media permissions and continuing without data loss.", MINT_BG, ACCENT_MINT),
    ("Self-Installing 1-Click Launcher", "Plug-and-Play start.bat", "Intelligent launcher verifies Python and Node.js, auto-installs missing npm and pip dependencies, configures .env, and boots both servers simultaneously.", PEACH_BG, ACCENT_PEACH),
    ("Future Scalability Roadmap", "Enterprise & Global Recruiting", "Next milestones: Custom company knowledge base simulation, ATS integrations (Greenhouse/Workday), and real-time multilingual roleplays.", LILAC_BG, ACCENT_LILAC)
]

for idx, (title, sub, desc, bg, acc) in enumerate(eng_boxes):
    col = idx % 2
    row = idx // 2
    x = Inches(0.8 + col * 6.0)
    y = Inches(1.8 + row * 2.5)

    box = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.6), Inches(2.2))
    box.fill.solid()
    box.fill.fore_color.rgb = bg
    box.line.color.rgb = acc
    box.line.width = Pt(1)

    tf = box.text_frame
    tf.margin_left = Inches(0.4)
    tf.margin_top = Inches(0.3)
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = TEXT_PRIMARY

    p2 = tf.add_paragraph()
    p2.text = sub
    p2.font.size = Pt(11)
    p2.font.bold = True
    p2.font.color.rgb = acc
    p2.space_before = Pt(2)

    p3 = tf.add_paragraph()
    p3.text = desc
    p3.font.size = Pt(11)
    p3.font.color.rgb = TEXT_MUTED
    p3.space_before = Pt(6)

output_path = "c:\\Users\\vaibh\\OneDrive\\Desktop\\HACKATHON\\AI_Job_Simulator_Presentation.pptx"
prs.save(output_path)
print(f"Presentation generated successfully at: {output_path}")
